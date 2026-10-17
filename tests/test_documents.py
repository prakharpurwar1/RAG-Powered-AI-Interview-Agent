"""Tests for Phase 2 document parsing, validation, and personalized question generation.

Run: python tests/test_documents.py
"""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docx
import pypdf

from interviewer.config import InterviewConfig
from interviewer.documents import (
    DocumentError,
    extract_document_text,
    extract_text_from_docx,
    extract_text_from_pdf,
    validate_text_document,
)
from interviewer.engine import InterviewEngine
from interviewer.llm import MockLLM
from interviewer.questions import generate_personalized_questions, questions_for


def _create_sample_pdf(text: str) -> bytes:
    """Helper to create an in-memory PDF with text stream and Helvetica font."""
    content_bytes = text.encode("latin1")
    stream_data = b"BT /F1 14 Tf 50 700 Td (" + content_bytes + b") Tj ET"
    pdf = (
        b"%PDF-1.4\n"
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n"
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >> >>\nendobj\n"
        b"4 0 obj\n<< /Length " + str(len(stream_data)).encode("latin1") + b" >>\nstream\n"
        + stream_data + b"\nendstream\nendobj\n"
        b"trailer\n<< /Size 5 /Root 1 0 R >>\nstartxref\n0\n%%EOF"
    )
    return pdf


def _create_sample_docx(paragraphs: list, table_data: list = None) -> bytes:
    """Helper to create an in-memory DOCX with paragraphs and tables."""
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                table.rows[r_idx].cells[c_idx].text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Document Extraction & Validation Tests
# ---------------------------------------------------------------------------
def test_pdf_extraction_success():
    sample_text = "Alice Smith - Senior Python Engineer with Distributed Systems Experience"
    pdf_bytes = _create_sample_pdf(sample_text)
    text = extract_text_from_pdf(pdf_bytes, filename="resume.pdf")
    assert sample_text in text or "Alice Smith" in text


def test_docx_extraction_success():
    paragraphs = [
        "Bob Johnson",
        "Staff Software Engineer at TechCorp",
        "Led architecture of high-throughput payment pipeline processing 10k TPS.",
    ]
    tables = [
        ["Skill", "Proficiency"],
        ["Python", "Expert"],
        ["FastAPI", "Advanced"],
    ]
    docx_bytes = _create_sample_docx(paragraphs, table_data=tables)
    text = extract_text_from_docx(docx_bytes, filename="resume.docx")
    assert "Bob Johnson" in text
    assert "TechCorp" in text
    assert "Python | Expert" in text


def test_validation_unsupported_file_extension():
    try:
        extract_document_text("resume.png", b"fake image bytes")
        assert False, "Should raise DocumentError for .png"
    except DocumentError as e:
        assert "Unsupported file format '.png'" in str(e)


def test_validation_empty_file():
    try:
        extract_document_text("empty.pdf", b"")
        assert False, "Should raise DocumentError for empty bytes"
    except DocumentError as e:
        assert "empty (0 bytes)" in str(e)


def test_validation_file_size_exceeded():
    oversized = b"x" * (11 * 1024 * 1024)  # 11MB
    try:
        extract_document_text("large.pdf", oversized)
        assert False, "Should raise DocumentError for oversized file"
    except DocumentError as e:
        assert "exceeds maximum allowed size" in str(e)


def test_validation_corrupted_pdf():
    try:
        extract_document_text("corrupted.pdf", b"%PDF-invalid-garbled-data")
        assert False, "Should raise DocumentError for corrupted PDF"
    except DocumentError as e:
        assert "Unable to read PDF file" in str(e) or "corrupted" in str(e)


def test_validation_corrupted_docx():
    try:
        extract_document_text("corrupted.docx", b"not-a-valid-zip-file-stream")
        assert False, "Should raise DocumentError for corrupted DOCX"
    except DocumentError as e:
        assert "Unable to read Word document" in str(e)


def test_validation_empty_scanned_pdf():
    # PDF with blank page and no text
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    try:
        extract_text_from_pdf(buf.getvalue(), filename="scanned.pdf")
        assert False, "Should raise DocumentError for blank / unreadable text"
    except DocumentError as e:
        assert "No readable text could be extracted" in str(e)


def test_validate_text_document():
    # Valid text
    cleaned = validate_text_document("   We are looking for a Senior Backend Engineer to build APIs.   ")
    assert "Senior Backend Engineer" in cleaned

    # Too brief
    try:
        validate_text_document("Too short", min_length=20)
        assert False, "Should reject too brief text"
    except DocumentError as e:
        assert "too brief" in str(e)

    # Empty
    try:
        validate_text_document("   ", min_length=20)
        assert False, "Should reject empty text"
    except DocumentError as e:
        assert "cannot be empty" in str(e)


# ---------------------------------------------------------------------------
# Personalized Question Generation & Grounding Tests
# ---------------------------------------------------------------------------
def test_personalized_questions_with_resume_and_jd():
    cfg = InterviewConfig(
        role="Backend Engineer",
        level="senior",
        interview_type="technical",
        max_questions=4,
        resume_text="Owned Postgres sharding and Redis caching at Acme Corp.",
        job_description="Seeking Senior Engineer with high-scale database and caching background.",
    )
    llm = MockLLM()
    questions = generate_personalized_questions(cfg, llm)
    assert len(questions) >= 3
    # Check that questions have required fields
    for q in questions:
        assert q.id
        assert q.competency
        assert q.text
        assert q.followup_hint


def test_personalized_questions_fallback_when_no_docs():
    cfg = InterviewConfig(
        role="Frontend Engineer",
        level="mid",
        interview_type="behavioral",
        resume_text="",
        job_description="",
    )
    questions = generate_personalized_questions(cfg, MockLLM())
    # Should cleanly match standard question bank
    bank = questions_for("behavioral")
    assert len(questions) == len(bank)
    assert questions[0].id == bank[0].id


def test_engine_initializes_with_grounded_context():
    cfg = InterviewConfig(
        role="Full Stack Engineer",
        level="senior",
        interview_type="mixed",
        candidate_name="Priya",
        resume_text="Architected React + FastAPI microservices at CloudScale.",
        job_description="Build modern microservices and web interfaces with React and Python.",
    )
    eng = InterviewEngine(cfg, MockLLM())
    assert len(eng.questions) > 0

    # Ensure system prompt embeds the candidate's resume, JD, and grounding rule
    sys_prompt = eng._system_prompt()
    assert "CloudScale" in sys_prompt
    assert "GROUNDING RULE" in sys_prompt
    assert "Never invent experience" in sys_prompt

    # Test full interview run works smoothly
    greeting = eng.start()
    assert len(greeting) > 0
    resp = eng.respond("I designed the event schema and implemented caching.")
    assert "message" in resp


if __name__ == "__main__":
    fns = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for fn in fns:
        fn()
        print(f"  ok  {fn.__name__}")
    print(f"\n{len(fns)} tests passed.")
